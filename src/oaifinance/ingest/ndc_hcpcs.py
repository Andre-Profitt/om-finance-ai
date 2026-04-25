"""CMS NDC-HCPCS crosswalk ingestion.

Same dual-path design as cms_asp: sample by default, live fetch via the
CMS landing page when invoked with `live=True`. Live merge filters to the
HCPCS we have specialty tags for so the downstream schema is stable.
"""

from __future__ import annotations

import io
import logging
import re
import zipfile
from typing import Any

import polars as pl
import requests

from oaifinance.config import BRONZE_DIR, SAMPLES_DIR

LOG = logging.getLogger(__name__)

CMS_LANDING = (
    "https://www.cms.gov/medicare/medicare-part-b-drug-average-sales-price/asp-pricing-files"
)
CMS_HOST = "https://www.cms.gov"


def load_sample() -> pl.DataFrame:
    return pl.read_csv(SAMPLES_DIR / "ndc_hcpcs_sample.csv")


def _resolve_latest_crosswalk_zip() -> str:
    r = requests.get(CMS_LANDING, timeout=30)
    r.raise_for_status()
    hrefs = re.findall(r'href="(/files/zip/[^"]+\.zip)"', r.text)
    for href in hrefs:
        if "ndc-hcpcs-crosswalk" in href.lower():
            return CMS_HOST + href
    raise RuntimeError("could not resolve CMS NDC-HCPCS crosswalk ZIP URL")


def _extract_xlsx(zip_bytes: bytes) -> bytes:
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        candidates = [n for n in zf.namelist() if n.lower().endswith((".xlsx", ".xls"))]
        if not candidates:
            raise RuntimeError("no xlsx in CMS crosswalk zip")
        # prefer non-addendum files; usually a single sheet
        ranked = sorted(candidates, key=lambda n: (0 if "addendum" not in n.lower() else 1, len(n)))
        return zf.read(ranked[0])


def _parse_crosswalk_workbook(xlsx_bytes: bytes, target_hcpcs: set[str]) -> pl.DataFrame:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=True, read_only=True)
    ws = wb.active
    if ws is None:
        raise RuntimeError("crosswalk workbook has no active sheet")

    rows_cache: list[tuple[Any, ...]] = []
    header_idx: int | None = None
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        rows_cache.append(row)
        if header_idx is None:
            cells_upper = [str(c).upper() if c is not None else "" for c in row]
            if any("HCPCS" in c for c in cells_upper) and any("NDC" in c for c in cells_upper):
                header_idx = i
        if header_idx is not None and i > header_idx + 1:
            break

    if header_idx is None:
        raise RuntimeError("could not find header row with HCPCS + NDC")

    for row in ws.iter_rows(values_only=True, min_row=len(rows_cache) + 1):
        rows_cache.append(row)

    header = [str(c) if c is not None else "" for c in rows_cache[header_idx]]

    def _find(*names: str) -> int | None:
        for n in names:
            for j, h in enumerate(header):
                if n.lower() in h.lower():
                    return j
        return None

    ndc_col = _find("NDC")
    hcpcs_col = _find("HCPCS Code", "HCPCS")
    label_col = _find("Labeler", "Manufacturer")
    brand_col = _find("Brand", "Trade", "Product Name")
    pkg_col = _find("Package Size", "Package", "Pkg Size")
    units_col = _find("Billable Units", "Units per Package", "Units")

    if ndc_col is None or hcpcs_col is None:
        raise RuntimeError(f"crosswalk required columns missing (ndc={ndc_col}, hcpcs={hcpcs_col})")

    out = []
    for row in rows_cache[header_idx + 1 :]:
        if not row:
            continue
        hcpcs = row[hcpcs_col]
        ndc = row[ndc_col]
        if hcpcs is None or ndc is None:
            continue
        hcpcs_str = str(hcpcs).strip()
        if hcpcs_str not in target_hcpcs:
            continue
        ndc_str = str(ndc).strip()
        labeler = (
            str(row[label_col]).strip()
            if label_col is not None and row[label_col] is not None
            else ""
        )
        brand = (
            str(row[brand_col]).strip()
            if brand_col is not None and row[brand_col] is not None
            else ""
        )
        try:
            pkg = (
                int(float(row[pkg_col])) if pkg_col is not None and row[pkg_col] is not None else 1
            )
        except (TypeError, ValueError):
            pkg = 1
        try:
            units = (
                int(float(row[units_col]))
                if units_col is not None and row[units_col] is not None
                else 1
            )
        except (TypeError, ValueError):
            units = 1
        out.append(
            {
                "ndc_code": ndc_str,
                "hcpcs_code": hcpcs_str,
                "labeler_name": labeler,
                "brand_name": brand,
                "package_size_mg": pkg,
                "billed_units_per_package": units,
            }
        )
    if not out:
        raise RuntimeError("no target crosswalk rows parsed")
    return pl.DataFrame(out)


def _live_ingest() -> pl.DataFrame:
    sample = load_sample()
    target = set(sample["hcpcs_code"].to_list())

    cross_url = _resolve_latest_crosswalk_zip()
    LOG.info("CMS NDC-HCPCS crosswalk ZIP resolved: %s", cross_url)

    r = requests.get(cross_url, timeout=120)
    r.raise_for_status()
    xlsx = _extract_xlsx(r.content)
    live = _parse_crosswalk_workbook(xlsx, target)

    if len(live) < len(target):
        raise RuntimeError(
            f"crosswalk live ingest produced only {len(live)} rows — "
            "schema or coverage drift; falling back to sample"
        )
    return live


def ingest(live: bool = False) -> pl.DataFrame:
    if live:
        try:
            df = _live_ingest()
            LOG.info("CMS NDC-HCPCS live ingest: %d rows", len(df))
        except Exception as exc:
            LOG.warning("Live CMS NDC-HCPCS fetch failed (%s); falling back to sample.", exc)
            df = load_sample()
    else:
        df = load_sample()

    BRONZE_DIR.mkdir(parents=True, exist_ok=True)
    df.write_parquet(BRONZE_DIR / "ndc_hcpcs_raw.parquet")
    return df


if __name__ == "__main__":
    df = ingest()
    print(f"ingested {len(df)} NDC-HCPCS rows")
    print(df)
