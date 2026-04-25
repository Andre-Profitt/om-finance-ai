"""CMS ASP Part B pricing file ingestion.

Two paths:

  * Sample (default): committed CSV with curated HCPCS for oncology +
    multispecialty. Reproducible and offline.
  * Live (`--live`): scrapes the CMS landing page, downloads the most
    recent quarter's payment-limit ZIP, parses the Excel, and merges
    against the sample's specialty / biosimilar tags so the schema is
    stable. Falls back to the sample on any error (no network in CI by
    default; opt-in via env or CLI flag).

CMS page: https://www.cms.gov/medicare/medicare-part-b-drug-average-sales-price/asp-pricing-files
"""

from __future__ import annotations

import io
import logging
import re
import zipfile
from typing import Any

import polars as pl
import requests

from oaifinance.config import BRONZE_DIR, SAMPLES_DIR

LOG = logging.getLogger(__name__)

CMS_LANDING = (
    "https://www.cms.gov/medicare/medicare-part-b-drug-average-sales-price/asp-pricing-files"
)
CMS_HOST = "https://www.cms.gov"


def load_sample() -> pl.DataFrame:
    return pl.read_csv(SAMPLES_DIR / "cms_asp_sample.csv", try_parse_dates=True)


def _resolve_latest_asp_zip() -> str:
    """Find the most recent quarter's ASP / payment-limit ZIP URL on the
    CMS landing page. Excludes crosswalk and NOC files.
    """
    r = requests.get(CMS_LANDING, timeout=30)
    r.raise_for_status()
    hrefs = re.findall(r'href="(/files/zip/[^"]+\.zip)"', r.text)
    for href in hrefs:
        lower = href.lower()
        if "crosswalk" in lower or "/noc-" in lower or "noc-pricing" in lower:
            continue
        if "asp-pricing" in lower or "payment-limit-files" in lower:
            return CMS_HOST + href
    raise RuntimeError("could not resolve CMS ASP ZIP URL from landing page")


def _quarter_effective_date(zip_url: str) -> str:
    """Best-effort effective date from the URL slug (e.g. april-2026 -> 2026-04-01)."""
    slug = zip_url.lower()
    months = {
        "january": "01",
        "april": "04",
        "july": "07",
        "october": "10",
    }
    year_match = re.search(r"-(\d{4})-", slug)
    year = year_match.group(1) if year_match else None
    for name, mm in months.items():
        if name in slug and year:
            return f"{year}-{mm}-01"
    return "2026-04-01"


def _extract_xlsx(zip_bytes: bytes, prefer_substr: str = "asp") -> bytes:
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        candidates = [n for n in zf.namelist() if n.lower().endswith((".xlsx", ".xls"))]
        if not candidates:
            raise RuntimeError("no xlsx in CMS zip")
        # prefer files whose name hints at ASP / payment-limit content
        ranked = sorted(
            candidates,
            key=lambda n: (
                0 if prefer_substr in n.lower() else 1,
                0 if "addendum" not in n.lower() else 1,
                len(n),
            ),
        )
        return zf.read(ranked[0])


def _parse_asp_workbook(xlsx_bytes: bytes, target_hcpcs: set[str]) -> pl.DataFrame:
    """Parse a CMS ASP / payment-limit XLSX. Filters to target HCPCS list and
    returns rows in our schema (live prices, sample-derived specialty tags
    merged downstream)."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=True, read_only=True)
    ws = wb.active
    if ws is None:
        raise RuntimeError("workbook has no active sheet")

    header_idx: int | None = None
    rows_cache: list[tuple[Any, ...]] = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        rows_cache.append(row)
        if i > 30 and header_idx is None:
            break
        cells_upper = [str(c).upper() if c is not None else "" for c in row]
        if any("HCPCS" in c for c in cells_upper):
            header_idx = i
            # keep iterating to populate rows_cache as far as needed
    # finish iterating after header is found
    if header_idx is None:
        raise RuntimeError("could not find header row containing 'HCPCS'")

    # gather remaining rows
    for row in ws.iter_rows(values_only=True, min_row=len(rows_cache) + 1):
        rows_cache.append(row)

    header = [str(c) if c is not None else "" for c in rows_cache[header_idx]]

    def _find(*names: str) -> int | None:
        for n in names:
            for j, h in enumerate(header):
                if n.lower() in h.lower():
                    return j
        return None

    hcpcs_col = _find("HCPCS Code", "HCPCS")
    desc_col = _find("Short Description", "Description")
    dosage_col = _find("HCPCS Code Dosage", "Dosage")
    limit_col = _find("Payment Limit", "Limit")

    if hcpcs_col is None or limit_col is None:
        raise RuntimeError(f"required columns missing (hcpcs={hcpcs_col}, limit={limit_col})")

    out = []
    for row in rows_cache[header_idx + 1 :]:
        if not row or row[hcpcs_col] is None:
            continue
        hcpcs = str(row[hcpcs_col]).strip()
        if hcpcs not in target_hcpcs:
            continue
        dosage_raw = (
            str(row[dosage_col]) if dosage_col is not None and row[dosage_col] is not None else "1"
        )
        m = re.search(r"(\d+(?:\.\d+)?)", dosage_raw)
        dosage_mg = max(1, int(float(m.group(1)))) if m else 1
        try:
            limit = float(row[limit_col]) if row[limit_col] is not None else None
        except (TypeError, ValueError):
            continue
        if limit is None or limit <= 0:
            continue
        desc = (
            str(row[desc_col]).strip() if desc_col is not None and row[desc_col] is not None else ""
        )
        out.append(
            {
                "hcpcs_code": hcpcs,
                "short_description": desc,
                "dosage_per_unit_mg": dosage_mg,
                "payment_limit_per_unit": limit,
            }
        )
    if not out:
        raise RuntimeError("no target HCPCS rows parsed from workbook")
    return pl.DataFrame(out)


def _live_ingest() -> pl.DataFrame:
    """Resolve, download, parse, and merge live CMS ASP data with our tags."""
    sample = load_sample()
    target = set(sample["hcpcs_code"].to_list())

    asp_url = _resolve_latest_asp_zip()
    LOG.info("CMS ASP ZIP resolved: %s", asp_url)

    r = requests.get(asp_url, timeout=120)
    r.raise_for_status()
    xlsx_bytes = _extract_xlsx(r.content, prefer_substr="asp")
    live = _parse_asp_workbook(xlsx_bytes, target)

    eff_date = _quarter_effective_date(asp_url)

    # Merge live prices with sample-derived specialty + biosimilar tags so
    # downstream consumers see a stable schema regardless of source.
    tags = sample.select("hcpcs_code", "biosimilar_reference", "specialty").unique(
        subset=["hcpcs_code"]
    )
    merged = (
        live.join(tags, on="hcpcs_code", how="inner")
        .with_columns(
            pl.lit(eff_date).str.strptime(pl.Date, "%Y-%m-%d").alias("effective_date"),
        )
        .select(
            "hcpcs_code",
            "short_description",
            "dosage_per_unit_mg",
            "payment_limit_per_unit",
            "effective_date",
            "biosimilar_reference",
            "specialty",
        )
    )

    if len(merged) < len(target) // 2:
        raise RuntimeError(
            f"live ingest produced only {len(merged)} of {len(target)} target HCPCS — "
            "schema or coverage drift; falling back to sample"
        )
    return merged


def ingest(live: bool = False) -> pl.DataFrame:
    if live:
        try:
            df = _live_ingest()
            LOG.info("CMS ASP live ingest: %d rows", len(df))
        except Exception as exc:  # network, parse, schema drift
            LOG.warning("Live CMS ASP fetch failed (%s); falling back to sample.", exc)
            df = load_sample()
    else:
        df = load_sample()

    BRONZE_DIR.mkdir(parents=True, exist_ok=True)
    df.write_parquet(BRONZE_DIR / "asp_raw.parquet")
    return df


if __name__ == "__main__":
    df = ingest()
    print(f"ingested {len(df)} ASP rows")
    print(df)
